from io import BytesIO

import pytest
from docx import Document as DocxDocument
from pypdf import PdfWriter

from rosevear_ai_hub.knowledge.ingestion import (
    KnowledgeIngestionError,
    KnowledgeIngestionService,
)


def service(tmp_path) -> KnowledgeIngestionService:
    return KnowledgeIngestionService(
        tmp_path / "knowledge",
        max_upload_bytes=2 * 1024 * 1024,
        max_expanded_docx_bytes=8 * 1024 * 1024,
    )


def test_text_ingestion_hashes_extracts_and_sanitizes_filename(tmp_path) -> None:
    prepared = service(tmp_path).prepare(
        filename="../../notes.md",
        content_type="text/markdown",
        data=b"# Workshop\nLocal notes.",
    )

    assert prepared.filename == "notes.md"
    assert prepared.mime_type == "text/markdown"
    assert prepared.extracted_text == "# Workshop\nLocal notes."
    assert len(prepared.content_hash) == 64
    assert prepared.metadata["character_count"] == len(prepared.extracted_text)


def test_docx_ingestion_extracts_paragraphs_and_tables(tmp_path) -> None:
    document = DocxDocument()
    document.add_paragraph("Workshop setup")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Tool"
    table.cell(0, 1).text = "Lathe"

    buffer = BytesIO()
    document.save(buffer)

    prepared = service(tmp_path).prepare(
        filename="workshop.docx",
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        data=buffer.getvalue(),
    )

    assert "Workshop setup" in prepared.extracted_text
    assert "Tool\tLathe" in prepared.extracted_text
    assert prepared.metadata["table_count"] == 1


def test_pdf_ingestion_accepts_valid_pdf_and_counts_pages(tmp_path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = BytesIO()
    writer.write(buffer)

    prepared = service(tmp_path).prepare(
        filename="scan.pdf",
        content_type="application/pdf",
        data=buffer.getvalue(),
    )

    assert prepared.mime_type == "application/pdf"
    assert prepared.metadata["page_count"] == 1
    assert prepared.metadata["text_extracted"] is False


def test_unsupported_and_oversized_files_are_rejected(tmp_path) -> None:
    ingestion = KnowledgeIngestionService(
        tmp_path / "knowledge",
        max_upload_bytes=4,
        max_expanded_docx_bytes=100,
    )

    with pytest.raises(KnowledgeIngestionError, match="Unsupported file type"):
        ingestion.prepare(
            filename="archive.zip",
            content_type="application/zip",
            data=b"zip",
        )

    with pytest.raises(KnowledgeIngestionError, match="upload limit"):
        ingestion.prepare(
            filename="notes.txt",
            content_type="text/plain",
            data=b"12345",
        )


def test_storage_is_content_addressed(tmp_path) -> None:
    ingestion = service(tmp_path)
    data = b"same bytes"
    prepared = ingestion.prepare(
        filename="notes.txt",
        content_type="text/plain",
        data=data,
    )

    first = ingestion.store(prepared, data)
    second = ingestion.store(prepared, data)

    assert first == second
    stored = ingestion.storage_root / first
    assert stored.read_bytes() == data
