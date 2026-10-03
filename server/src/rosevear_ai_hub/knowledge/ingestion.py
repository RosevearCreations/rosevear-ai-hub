"""Safe local file validation, extraction, hashing, and storage."""

from __future__ import annotations

import hashlib
import mimetypes
import tempfile
import zipfile
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from docx import Document as DocxDocument
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {".pdf", ".txt", ".md", ".docx"}
MIME_BY_EXTENSION = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
}


class KnowledgeIngestionError(ValueError):
    """Base validation/extraction error with an HTTP-friendly status code."""

    def __init__(self, message: str, *, status_code: int = 422) -> None:
        super().__init__(message)
        self.status_code = status_code


@dataclass(frozen=True)
class PreparedDocument:
    """Validated document ready for metadata persistence and local storage."""

    filename: str
    extension: str
    mime_type: str
    content_hash: str
    size_bytes: int
    extracted_text: str
    metadata: dict[str, int | str | bool | None]


class KnowledgeIngestionService:
    """Validate, extract, hash, and store supported local documents."""

    def __init__(
        self,
        storage_root: Path,
        *,
        max_upload_bytes: int,
        max_expanded_docx_bytes: int,
    ) -> None:
        self.storage_root = storage_root
        self.max_upload_bytes = max_upload_bytes
        self.max_expanded_docx_bytes = max_expanded_docx_bytes

    def prepare(
        self,
        *,
        filename: str,
        content_type: str | None,
        data: bytes,
    ) -> PreparedDocument:
        safe_filename = Path(filename.replace("\\", "/")).name.strip()
        if not safe_filename or safe_filename in {".", ".."}:
            raise KnowledgeIngestionError("A valid filename is required.", status_code=400)

        extension = Path(safe_filename).suffix.lower()
        if extension not in SUPPORTED_EXTENSIONS:
            raise KnowledgeIngestionError(
                "Unsupported file type. Allowed types: PDF, TXT, Markdown, DOCX.",
                status_code=415,
            )

        if not data:
            raise KnowledgeIngestionError("The uploaded file is empty.", status_code=400)

        if len(data) > self.max_upload_bytes:
            raise KnowledgeIngestionError(
                f"File exceeds the {self.max_upload_bytes // (1024 * 1024)} MB upload limit.",
                status_code=413,
            )

        expected_mime = MIME_BY_EXTENSION[extension]
        normalized_content_type = (content_type or "").split(";", 1)[0].strip().lower()
        allowed_content_types = {
            "",
            "application/octet-stream",
            expected_mime,
            mimetypes.guess_type(safe_filename)[0] or expected_mime,
        }
        if extension == ".md":
            allowed_content_types.add("text/plain")
        if normalized_content_type not in allowed_content_types:
            raise KnowledgeIngestionError(
                f"Content type does not match the {extension} file extension.",
                status_code=415,
            )

        content_hash = hashlib.sha256(data).hexdigest()
        extracted_text, metadata = self._extract(extension, data)

        return PreparedDocument(
            filename=safe_filename,
            extension=extension,
            mime_type=expected_mime,
            content_hash=content_hash,
            size_bytes=len(data),
            extracted_text=extracted_text,
            metadata=metadata,
        )

    def store(self, prepared: PreparedDocument, data: bytes) -> str:
        """Persist original bytes under a content-addressed relative path."""

        originals = self.storage_root / "originals"
        originals.mkdir(parents=True, exist_ok=True)
        target = originals / f"{prepared.content_hash}{prepared.extension}"

        if not target.exists():
            with tempfile.NamedTemporaryFile(
                mode="wb",
                prefix=f".{prepared.content_hash}.",
                suffix=".tmp",
                dir=originals,
                delete=False,
            ) as temporary:
                temporary.write(data)
                temporary.flush()
                temporary_path = Path(temporary.name)

            temporary_path.replace(target)

        return target.relative_to(self.storage_root).as_posix()

    def _extract(
        self,
        extension: str,
        data: bytes,
    ) -> tuple[str, dict[str, int | str | bool | None]]:
        if extension in {".txt", ".md"}:
            try:
                text = data.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise KnowledgeIngestionError(
                    "Text and Markdown files must use UTF-8 encoding."
                ) from exc
            return text, {
                "format": extension.lstrip("."),
                "page_count": None,
                "character_count": len(text),
            }

        if extension == ".pdf":
            if not data.startswith(b"%PDF-"):
                raise KnowledgeIngestionError("The uploaded PDF signature is invalid.")

            try:
                reader = PdfReader(BytesIO(data))
                if reader.is_encrypted:
                    raise KnowledgeIngestionError(
                        "Encrypted PDFs are not supported in local knowledge ingestion."
                    )
                page_text = [(page.extract_text() or "") for page in reader.pages]
            except KnowledgeIngestionError:
                raise
            except Exception as exc:
                raise KnowledgeIngestionError("The PDF could not be parsed safely.") from exc

            text = "\n\n".join(page_text)
            return text, {
                "format": "pdf",
                "page_count": len(reader.pages),
                "character_count": len(text),
                "text_extracted": bool(text.strip()),
            }

        if extension == ".docx":
            self._validate_docx_archive(data)
            try:
                document = DocxDocument(BytesIO(data))
            except Exception as exc:
                raise KnowledgeIngestionError("The DOCX file could not be parsed safely.") from exc

            blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        blocks.append("\t".join(cells))

            text = "\n".join(blocks)
            return text, {
                "format": "docx",
                "page_count": None,
                "character_count": len(text),
                "paragraph_count": len(document.paragraphs),
                "table_count": len(document.tables),
            }

        raise KnowledgeIngestionError("Unsupported file type.", status_code=415)

    def _validate_docx_archive(self, data: bytes) -> None:
        try:
            with zipfile.ZipFile(BytesIO(data)) as archive:
                names = set(archive.namelist())
                if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                    raise KnowledgeIngestionError("The uploaded DOCX structure is invalid.")

                expanded_bytes = sum(item.file_size for item in archive.infolist())
                if expanded_bytes > self.max_expanded_docx_bytes:
                    raise KnowledgeIngestionError(
                        "The DOCX expands beyond the safe processing limit.",
                        status_code=413,
                    )
        except KnowledgeIngestionError:
            raise
        except (zipfile.BadZipFile, OSError) as exc:
            raise KnowledgeIngestionError("The uploaded DOCX archive is invalid.") from exc
